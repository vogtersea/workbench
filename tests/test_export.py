"""报告导出（W-01）纯函数用例：只测 app.services.export，不依赖服务端与数据库。

依赖：pytest、python-docx（requirements.txt）。运行：pytest tests/test_export.py
"""

from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

# 便于单独运行本文件（tests/conftest.py 由另一代理维护，这里不重复造它）
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.export import (
    BODY_FONT,
    H1_FONT,
    H2_FONT,
    INDENT_PT,
    LINE_PT,
    TITLE_FONT,
    build_docx,
)

DOCX_MAGIC = b"PK\x03\x04"


def _load(data: bytes) -> Document:
    """把字节流重新读成 Document，验证产物可被 Word/python-docx 打开。"""
    return Document(io.BytesIO(data))


def _xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _east_asia(run) -> str | None:
    rpr = run._element.find(qn("w:rPr"))
    fonts = None if rpr is None else rpr.find(qn("w:rFonts"))
    return None if fonts is None else fonts.get(qn("w:eastAsia"))


def test_无一级标题时回退到_fallback_title():
    data = build_docx("正文第一段。\n正文第二段。", "周报：2026-W42")

    assert data.startswith(DOCX_MAGIC)
    doc = _load(data)
    assert doc.paragraphs[0].text == "周报：2026-W42"
    assert doc.paragraphs[0].runs[0].font.name == TITLE_FONT
    assert doc.paragraphs[0].runs[0].font.size == Pt(22)
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert [p.text for p in doc.paragraphs[1:]] == ["正文第一段。", "正文第二段。"]


def test_首个一级标题提升为大标题且不在正文重复():
    data = build_docx("# 本周工作小结\n\n正文。", "回退标题")

    doc = _load(data)
    assert [p.text for p in doc.paragraphs] == ["本周工作小结", "正文。"]


def test_空正文仍生成可打开的合法_docx():
    data = build_docx("")
    assert data.startswith(DOCX_MAGIC)

    doc = _load(data)
    assert all(not p.text.strip() for p in doc.paragraphs)
    section = doc.sections[0]
    assert round(section.top_margin.cm, 2) == 3.5
    assert round(section.bottom_margin.cm, 2) == 3.3

    titled = _load(build_docx("", "只有标题的报告"))
    assert [p.text for p in titled.paragraphs] == ["只有标题的报告"]


def test_加粗与链接标记只保留纯文本():
    markdown = "本周完成 **关键改造**，详见 [方案文档](https://example.com/plan)。"
    data = build_docx(markdown, "回退标题")

    paragraph = _load(data).paragraphs[-1]
    run = paragraph.runs[0]
    assert paragraph.text == "本周完成 关键改造，详见 方案文档。"
    assert "**" not in paragraph.text
    assert "](" not in paragraph.text
    assert run.font.name == BODY_FONT
    assert _east_asia(run) == BODY_FONT
    assert run.font.size == Pt(16)
    assert paragraph.paragraph_format.first_line_indent == Pt(INDENT_PT)
    assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert paragraph.paragraph_format.line_spacing == Pt(LINE_PT)


def test_超长单段不拆分不丢字():
    long_line = "这是一条用于验证超长段落不会被截断或拆分的正文，" * 300
    assert len(long_line) > 6000

    doc = _load(build_docx(long_line, "超长段落"))
    assert len(doc.paragraphs) == 2
    assert doc.paragraphs[-1].text == long_line
    assert doc.paragraphs[-1].paragraph_format.first_line_indent == Pt(INDENT_PT)


def test_小节按级别映射黑体楷体仿宋并自动编号():
    markdown = "## 一级小节\n\n### 二级小节\n\n#### 三级小节\n\n正文段落。"
    doc = _load(build_docx(markdown, "标题"))

    # v1.0.4 起按公文层级自动编号：一、/（一）/1.
    assert [p.text for p in doc.paragraphs[1:]] == [
        "一、一级小节",
        "（一）二级小节",
        "1.三级小节",
        "正文段落。",
    ]
    fonts = [p.runs[0].font.name for p in doc.paragraphs[1:]]
    assert fonts == [H1_FONT, H2_FONT, BODY_FONT, BODY_FONT]


def test_document_xml_写入公文字体与固定行距():
    data = build_docx("# 标题\n\n## 小节\n\n正文段落。")
    xml = _xml(data)

    assert f'w:eastAsia="{BODY_FONT}"' in xml
    assert f'w:eastAsia="{TITLE_FONT}"' in xml
    assert f'w:eastAsia="{H1_FONT}"' in xml
    assert 'w:firstLine="640"' in xml  # 32pt = 640 twips
    assert 'w:line="590"' in xml  # 29.5pt = 590 twips
    spacing = re.search(r"<w:spacing[^>]*/>", xml)
    assert spacing is not None
    assert 'w:line="590"' in spacing.group(0)
    assert 'w:lineRule="exact"' in spacing.group(0)


def test_页面设置符合公文页边距():
    section = _load(build_docx("# 标题\n\n正文。")).sections[0]
    assert round(section.top_margin.cm, 2) == 3.5
    assert round(section.bottom_margin.cm, 2) == 3.3
    assert round(section.left_margin.cm, 2) == 2.7
    assert round(section.right_margin.cm, 2) == 2.7
