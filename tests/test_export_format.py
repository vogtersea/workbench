"""docx 导出格式：逐项对照《公文写作格式规范》。

断言的是 docx 内部 XML 的真实属性（页边距、字体槽位、行距、缩进字符数），
不靠肉眼，也不靠 python-docx 的默认值。
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

from app.services import export

SAMPLE = """# 关于本周工作情况的报告

## 概述

本周完成三项重点工作。

## 完成情况

### 网讯稿件

初稿已交审。

## 问题与风险

暂无。

## 下阶段计划

继续推进。
"""


def _doc(markdown: str = SAMPLE):
    return Document(io.BytesIO(export.build_docx(markdown)))


def _font(run) -> str:
    rfonts = run._element.rPr.find(qn("w:rFonts"))
    return rfonts.get(qn("w:eastAsia"))


def _ind(paragraph, attr: str):
    ind = paragraph._p.pPr.find(qn("w:ind"))
    return None if ind is None else ind.get(qn(attr))


def test_页边距与页眉页脚距():
    section = _doc().sections[0]
    assert section.top_margin.cm == pytest.approx(3.5, abs=0.01)
    assert section.bottom_margin.cm == pytest.approx(3.3, abs=0.01)
    assert section.left_margin.cm == pytest.approx(2.7, abs=0.01)
    assert section.right_margin.cm == pytest.approx(2.7, abs=0.01)
    assert section.header_distance.cm == pytest.approx(1.5, abs=0.01)
    assert section.footer_distance.cm == pytest.approx(1.75, abs=0.01)


def test_大标题二号小标宋居中不加粗():
    para = _doc().paragraphs[0]
    assert para.text == "关于本周工作情况的报告"
    assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    run = para.runs[0]
    assert _font(run) == "方正小标宋简体"
    assert run.font.size.pt == 22
    assert run.font.bold is False
    assert _ind(para, "w:firstLineChars") in ("0", None)


def test_四级标题的字体与编号():
    doc = _doc(
        "# 标题\n\n## 一级甲\n\n### 二级甲\n\n#### 三级甲\n\n##### 四级甲\n"
    )
    texts = [p.text for p in doc.paragraphs]
    assert "一、一级甲" in texts
    assert "（一）二级甲" in texts
    assert "1.三级甲" in texts
    assert "（1）四级甲" in texts

    by_text = {p.text: p for p in doc.paragraphs}
    assert _font(by_text["一、一级甲"].runs[0]) == "黑体"
    assert _font(by_text["（一）二级甲"].runs[0]) == "楷体_GB2312"
    assert _font(by_text["1.三级甲"].runs[0]) == "仿宋_GB2312"
    assert _font(by_text["（1）四级甲"].runs[0]) == "仿宋_GB2312"
    for text in ("一、一级甲", "（一）二级甲", "1.三级甲", "（1）四级甲"):
        assert by_text[text].runs[0].font.size.pt == 16
        assert by_text[text].runs[0].font.bold is False


def test_编号连续且子级重置():
    doc = _doc(
        "# 标题\n\n## 甲\n\n### 甲一\n\n### 甲二\n\n## 乙\n\n### 乙一\n"
    )
    texts = [p.text for p in doc.paragraphs]
    assert "一、甲" in texts
    assert "（一）甲一" in texts
    assert "（二）甲二" in texts
    assert "二、乙" in texts
    assert "（一）乙一" in texts  # 换了一级标题后二级重新从（一）开始


def test_自带编号不重复编号():
    doc = _doc("# 标题\n\n## 二、风险\n\n## 计划\n")
    texts = [p.text for p in doc.paragraphs]
    assert "二、风险" in texts
    assert "三、计划" in texts  # 同步计数器：接着自带的「二」往后排


def test_正文仿宋三号两端对齐固定行距首行缩进两字符():
    doc = _doc()
    body = [p for p in doc.paragraphs if p.text == "本周完成三项重点工作。"][0]
    run = body.runs[0]
    assert _font(run) == "仿宋_GB2312"
    assert run.font.size.pt == 16
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = body.paragraph_format
    assert pf.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert pf.line_spacing.pt == pytest.approx(29.5, abs=0.01)
    assert _ind(body, "w:firstLineChars") == "200"


def test_数字字体与汉字一致():
    doc = _doc("# 标题\n\n2026 年完成 12 项，占 60%。\n")
    body = [p for p in doc.paragraphs if "2026" in p.text][0]
    rfonts = body.runs[0]._element.rPr.find(qn("w:rFonts"))
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        assert rfonts.get(qn(attr)) == "仿宋_GB2312"


def test_全文没有任何加粗():
    doc = _doc("# 标题\n\n## **加粗标题**\n\n**加粗正文**\n")
    for para in doc.paragraphs:
        for run in para.runs:
            assert run.font.bold is False
    texts = [p.text for p in doc.paragraphs]
    assert "一、加粗标题" in texts  # 行内 ** 标记被剥掉
    assert "加粗正文" in texts


def test_与标题重复的首个一级标题被跳过():
    doc = _doc("# 关于本周情况的报告\n\n## 关于本周情况的报告\n\n## 概述\n\n正文。\n")
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "关于本周情况的报告"
    assert "关于本周情况的报告" not in texts[1:]
    assert "一、概述" in texts


def test_空标题不产生空段落():
    doc = _doc("#\n\n##\n\n正文\n")
    assert [p.text for p in doc.paragraphs] == ["正文"]
