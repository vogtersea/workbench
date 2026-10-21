"""报告导出：Markdown → docx（公文格式）。

排版规则（中文公文通用约定）：
  纸张 A4；页边距 上3.5 下3.3 左2.7 右2.7 cm；页眉 1.5 cm、页脚 1.75 cm
  行间距 固定值 29.5 磅；首行缩进 2 字符（w:firstLineChars=200，按字符而非磅值）
  正文 仿宋_GB2312 三号（16pt），两端对齐
  大标题 方正小标宋简体 二号（22pt），居中
  一级 黑体 三号「一、」；二级 楷体_GB2312 三号「（一）」；
  三级 仿宋 三号「1.」；四级 仿宋 三号「（1）」
  所有标题不加粗；数字与汉字同字体（ascii/hAnsi/eastAsia/cs 四槽同名字体）

注意：字体只写名称不嵌入，打开端缺 `仿宋_GB2312` 等字体会回退到系统默认字体。
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BODY_FONT = "仿宋_GB2312"
TITLE_FONT = "方正小标宋简体"
BODY_SIZE = 16          # 三号
TITLE_SIZE = 22         # 二号
LINE_PT = 29.5
INDENT_CHARS = 200      # w:firstLineChars 单位是百分之一字符

# 兼容旧调用方/测试的常量别名（值仍与 HEADING_STYLE 一致）
H1_FONT = "黑体"
H2_FONT = "楷体_GB2312"
INDENT_PT = BODY_SIZE * 2   # 首行缩进 2 字符 = 32 磅（同时写 firstLineChars）

#: 公文层级 -> (字体, 是否缩进)
HEADING_STYLE = {
    1: (("黑体", BODY_SIZE), True),
    2: (("楷体_GB2312", BODY_SIZE), True),
    3: (("仿宋_GB2312", BODY_SIZE), True),
    4: (("仿宋_GB2312", BODY_SIZE), True),
}

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_LIST_RE = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
_HR_RE = re.compile(r"^\s*(?:-{3,}|\*{3,})\s*$")

#: 已带编号的标题（模型有时自己写「一、」），识别后不再重复编号
_PREFIX = {
    1: re.compile(r"^([一二三四五六七八九十百]+)、"),
    2: re.compile(r"^（([一二三四五六七八九十百]+)）"),
    3: re.compile(r"^(\d+)[.、]"),
    4: re.compile(r"^（(\d+)）"),
}

_CN_DIGITS = "零一二三四五六七八九"


def _cn_number(n: int) -> str:
    """1 -> 一，10 -> 十，12 -> 十二，20 -> 二十（公文编号只用到百以内）。"""
    if n <= 0:
        return _CN_DIGITS[0]
    if n < 10:
        return _CN_DIGITS[n]
    if n < 20:
        return "十" + (_CN_DIGITS[n % 10] if n % 10 else "")
    if n < 100:
        return _CN_DIGITS[n // 10] + "十" + (_CN_DIGITS[n % 10] if n % 10 else "")
    return str(n)


def _cn_to_int(text: str) -> int:
    """中文数字转整数（够用即可：一 ~ 九十九）。"""
    if text.isdigit():
        return int(text)
    total, current = 0, 0
    for char in text:
        if char == "十":
            total += (current or 1) * 10
            current = 0
        elif char in _CN_DIGITS:
            current = _CN_DIGITS.index(char)
    return total + current


class _Numberer:
    """公文层级自动编号：一级「一、」，二级「（一）」，三级「1.」，四级「（1）」。

    子级出现时重置其下所有层级；若标题自带编号（模型自己写了「二、」），
    以自带编号为准并同步计数器，避免后文编号错位。
    """

    def __init__(self) -> None:
        self.counts = [0, 0, 0, 0]

    def take(self, level: int, text: str) -> str:
        existing = _PREFIX[level].match(text)
        if existing:
            self.counts[level - 1] = _cn_to_int(existing.group(1))
            self._reset_below(level)
            return text
        self.counts[level - 1] += 1
        self._reset_below(level)
        n = self.counts[level - 1]
        if level == 1:
            prefix = _cn_number(n) + "、"
        elif level == 2:
            prefix = "（" + _cn_number(n) + "）"
        elif level == 3:
            prefix = "%d." % n
        else:
            prefix = "（%d）" % n
        return prefix + text

    def _reset_below(self, level: int) -> None:
        for index in range(level, 4):
            self.counts[index] = 0


def _clean(text: str) -> str:
    """去掉行内 Markdown 标记，保留纯文本。"""
    text = _LINK_RE.sub(r"\1", text)
    text = _BOLD_RE.sub(r"\1", text)
    return text.replace("`", "").strip()


def _set_font(run, name: str, size: int) -> None:
    """四个字体槽位写同一个名字：数字与汉字同字体，不单独设置西文字体。"""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = False
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _set_indent(paragraph, chars: int) -> None:
    """首行缩进按「字符」设置（Word 里的 2 字符），同时给出磅值兜底。"""
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = ppr.makeelement(qn("w:ind"), {})
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), str(Pt(BODY_SIZE * chars / 100).twips))


def _format(paragraph, *, indent: bool, align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_PT)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    _set_indent(paragraph, INDENT_CHARS if indent else 0)
    paragraph.alignment = align


def _add(doc, text, font, size, *, indent=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = doc.add_paragraph()
    _format(paragraph, indent=indent, align=align)
    _set_font(paragraph.add_run(text), font, size)
    return paragraph


def _heading_level(md_level: int) -> int:
    """Markdown 标题层级 -> 公文层级。## 与 # 都算一级，往上逐级递减。"""
    if md_level <= 2:
        return 1
    if md_level == 3:
        return 2
    if md_level == 4:
        return 3
    return 4


def build_docx(markdown: str, fallback_title: str = "") -> bytes:
    """把报告 Markdown 渲染为 docx 字节流。"""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(3.5)
    section.bottom_margin = Cm(3.3)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)

    lines = (markdown or "").replace("\r\n", "\n").split("\n")

    # 首个一级标题提升为大标题，避免标题在正文里重复出现
    title = fallback_title
    start = 0
    for index, raw in enumerate(lines):
        if not raw.strip():
            continue
        if raw.lstrip().startswith("# "):
            title = _clean(raw.lstrip()[2:]) or fallback_title
            start = index + 1
        break

    if title:
        _add(doc, title, TITLE_FONT, TITLE_SIZE, align=WD_ALIGN_PARAGRAPH.CENTER)

    numberer = _Numberer()
    title_norm = re.sub(r"\s+", "", title or "")
    seen_heading = False
    for raw in lines[start:]:
        line = raw.rstrip()
        if not line.strip() or _HR_RE.match(line):
            continue
        stripped = line.lstrip()
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = _clean(stripped[level:])
            if not text:
                continue
            gongwen_level = _heading_level(level)
            # 首个一级标题若与文档标题相同（模型常把小节写成标题全文），跳过以免重复
            if not seen_heading and gongwen_level == 1 and re.sub(r"\s+", "", text) == title_norm:
                seen_heading = True
                continue
            seen_heading = True
            text = numberer.take(gongwen_level, text)
            (font, size), indent = HEADING_STYLE[gongwen_level]
            _add(doc, text, font, size, indent=indent)
            continue
        if stripped.startswith(">"):
            _add(doc, _clean(stripped[1:]), "楷体_GB2312", BODY_SIZE, indent=True)
            continue
        if _LIST_RE.match(line):
            _add(doc, _clean(_LIST_RE.sub("", line, count=1)), BODY_FONT, BODY_SIZE, indent=True)
            continue
        _add(doc, _clean(line), BODY_FONT, BODY_SIZE, indent=True)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
